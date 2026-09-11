"""Stretch pandoc's tables to the full text column.

Pandoc sizes docx columns from the *source* markdown widths, so a table whose
markdown was narrow renders at ~70% of the page and wraps text that had room to
breathe. Scale each table up to the usable width, keeping the column ratios.
"""
import sys
from docx import Document
from docx.shared import Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

path, usable_cm = sys.argv[1], float(sys.argv[2])

doc = Document(path)
fixed = 0
for t in doc.tables:
    widths = [c.width for c in t.columns]
    if any(w is None for w in widths):
        continue
    total = sum(w.cm for w in widths)
    if total <= 0:
        continue
    scale = usable_cm / total
    if abs(scale - 1.0) < 0.02:
        continue
    t.autofit = False
    # Table width as a percentage of the text column.
    tblPr = t._tbl.tblPr
    for old in tblPr.findall(qn('w:tblW')):
        tblPr.remove(old)
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:type'), 'pct')
    tblW.set(qn('w:w'), '5000')          # 5000 fiftieths of a percent = 100%
    tblPr.append(tblW)
    for col, w in zip(t.columns, widths):
        new = Cm(w.cm * scale)
        col.width = new
        for cell in col.cells:
            cell.width = new
    fixed += 1

# A row that breaks across a page leaves half a sentence stranded under the
# repeated header, which reads as a rendering bug rather than a long cell.
split_guarded = 0
for t in doc.tables:
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr()
        if trPr.find(qn('w:cantSplit')) is None:
            trPr.append(OxmlElement('w:cantSplit'))
            split_guarded += 1

doc.save(path)
print(f"  guarded {split_guarded} rows against splitting across pages")
print(f"  stretched {fixed} of {len(doc.tables)} tables to {usable_cm:.1f} cm")
