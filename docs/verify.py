"""Check the rendered documents, rather than assume they came out right.

Checks page count, page size, table overflow past the text column, headings
orphaned at the foot of a page, rupee glyphs surviving the conversion, and the
house rule that there are no em dashes anywhere.
"""
import os, sys
from docx import Document
from pypdf import PdfReader

SPECS = [("WHAT-I-BUILT", 17.2, None), ("DECISIONS", 17.8, 1)]
fail = []

for base, usable, max_pages in SPECS:
    doc, pdf = Document(f"{base}.docx"), PdfReader(f"{base}.pdf")
    pages, box = len(pdf.pages), pdf.pages[0].mediabox
    text = "\n".join(p.extract_text() for p in pdf.pages)

    over = []
    for i, t in enumerate(doc.tables):
        w = [c.width for c in t.columns]
        if all(x is not None for x in w):
            total = sum(x.cm for x in w)
            if total > usable + 0.05:
                over.append((i, round(total, 2)))

    heads = [p.text.strip() for p in doc.paragraphs
             if p.style.name.startswith("Heading") and p.text.strip()]
    orphans = []
    for pi, page in enumerate(pdf.pages):
        lines = [l.strip() for l in page.extract_text().split("\n") if l.strip()]
        if lines and any(lines[-1] == h or lines[-1].rstrip(":") == h for h in heads):
            orphans.append(pi + 1)

    print(f"{base}")
    print(f"  pages              {pages}" + (f" (max {max_pages})" if max_pages else ""))
    print(f"  page size          {float(box.width):.0f}x{float(box.height):.0f}pt (A4 = 595x842)")
    print(f"  file size          {os.path.getsize(base+'.pdf')/1024:.0f} KB pdf, "
          f"{os.path.getsize(base+'.docx')/1024:.0f} KB docx")
    print(f"  tables             {len(doc.tables)}, over the text column: {over or 'none'}")
    print(f"  headings           {len(heads)}, orphaned at a page foot: {orphans or 'none'}")
    print(f"  rupee glyphs       {text.count(chr(8377))}")
    print(f"  em dashes          {text.count(chr(8212))}")

    if abs(float(box.width) - 595) > 2: fail.append(f"{base}: not A4")
    if over:                            fail.append(f"{base}: table wider than the text column")
    if orphans:                         fail.append(f"{base}: heading orphaned at a page foot")
    if text.count(chr(8212)):           fail.append(f"{base}: em dash in the output")
    if not text.count(chr(8377)):       fail.append(f"{base}: rupee glyph lost in conversion")
    if max_pages and pages > max_pages: fail.append(f"{base}: {pages} pages, the brief asks for {max_pages}")
    print()

if fail:
    print("FAIL"); [print("  " + f) for f in fail]; sys.exit(1)
print("all checks pass")
