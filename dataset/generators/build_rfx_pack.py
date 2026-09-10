"""
The RFx pack, as four separate documents plus instructions.

A real procurement pack is never one file. It goes out as a scope document, a
priced line-item sheet, a questionnaire and a terms sheet, because different
people inside the vendor answer different parts: pre-sales does the scope,
inside sales does the pricing, compliance does the questionnaire, legal does
the terms.

Splitting them is also what creates the trap the prototype has to survive: the
questionnaire comes back on a different day, in a different format, from a
different person, than the prices. Nothing forces them to agree.

  01_Scope_of_Work.pdf            what is being bought, where, when
  02_Line_Items.xlsx              the 30 priced lines, the template to ignore
  03_Vendor_Questionnaire.xlsx    10 questions, 6 mandatory
  04_Commercial_Terms.pdf         tax, incoterm, payment, warranty, penalties
  05_Instructions_to_Bidders.pdf  how to respond, and the four pre-emptions
"""

from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                TableStyle)

from catalog import BUYER, RFX, LINES, QUESTIONNAIRE

OUT = Path(__file__).resolve().parents[1] / "out" / "00-rfx"

THIN = Side(style="thin", color="B0B0B5")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NAVY = "2F4F6F"
RED = "9C2A2A"

_styles = getSampleStyleSheet()
BODY = ParagraphStyle("body", parent=_styles["BodyText"], fontSize=9.5, leading=14)
SMALL = ParagraphStyle("small", parent=_styles["BodyText"], fontSize=8, leading=11.5)
CELL = ParagraphStyle("cell", parent=_styles["BodyText"], fontSize=7.6, leading=9.8)

written = []


def _stamp(title):
    def draw(canvas, _doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 6.6)
        canvas.setFillColorRGB(0.35, 0.35, 0.35)
        canvas.drawString(20 * mm, 9 * mm, f"{RFX['id']}  |  {title}")
        canvas.drawRightString(190 * mm, 9 * mm, f"Page {canvas.getPageNumber()}")
        canvas.restoreState()
    return draw


def _doc(path, title):
    return SimpleDocTemplate(str(path), pagesize=A4, topMargin=18 * mm,
                             bottomMargin=18 * mm, leftMargin=20 * mm,
                             rightMargin=20 * mm, title=f"{RFX['id']} {title}")


def _letterhead(title):
    return [
        Paragraph(f"<font size=13><b>{BUYER['legal_name']}</b></font>", BODY),
        Paragraph(f"<font size=7.5>{BUYER['hq']}  |  CIN {BUYER['cin']}  |  "
                  f"GSTIN {BUYER['gstin']}</font>", SMALL),
        Spacer(1, 4),
        Table([[""]], colWidths=[170 * mm], rowHeights=[1.4],
              style=TableStyle([("BACKGROUND", (0, 0), (-1, -1),
                                 colors.HexColor("#2F4F6F"))])),
        Spacer(1, 9),
        Paragraph(f"<font size=11><b>{title}</b></font>", BODY),
        Paragraph(f"<font size=8.5>{RFX['id']} - {RFX['title']}<br/>"
                  f"Issued {RFX['issued']:%d %B %Y}  |  Responses due "
                  f"{RFX['due']:%d %B %Y}, 17:00 IST</font>", SMALL),
        Spacer(1, 10),
    ]


def _sign_off():
    c = RFX["buyer_contact"]
    a = RFX["approver"]
    return [Spacer(1, 14),
            Paragraph(f"<b>{c['name']}</b><br/>{c['title']}<br/>"
                      f"{c['email']}  |  {c['phone']}", SMALL),
            Spacer(1, 6),
            Paragraph(f"<font size=7.5>Approving authority: {a['name']}, "
                      f"{a['title']}</font>", SMALL)]


# ===========================================================================
# 01  Scope of Work
# ===========================================================================

def scope_of_work():
    p = OUT / "01_Scope_of_Work.pdf"
    fl = _letterhead("Scope of Work")

    fl.append(Paragraph("<b>1. Background</b>", BODY))
    fl.append(Paragraph(
        f"{BUYER['short_name']} operates {BUYER['headcount']} staff across three "
        f"locations: {', '.join(BUYER['sites'])}. The current enduser fleet was "
        f"acquired in FY23 and reaches end of standard warranty during FY27. This "
        f"enquiry covers the refresh of that fleet together with the datacentre and "
        f"network hardware required to support it.", BODY))
    fl.append(Spacer(1, 7))

    fl.append(Paragraph("<b>2. Scope</b>", BODY))
    fl.append(Paragraph(
        f"Supply, delivery, installation and commissioning of {len(LINES)} line "
        f"items across four groups, as specified in <i>02_Line_Items.xlsx</i>. "
        f"Quantities in that sheet are firm.", BODY))
    fl.append(Spacer(1, 5))

    groups = {}
    for ln in LINES:
        groups.setdefault(ln["group"], []).append(ln)
    rows = [[Paragraph("<b>Group</b>", CELL), Paragraph("<b>Lines</b>", CELL),
             Paragraph("<b>Units</b>", CELL), Paragraph("<b>Covers</b>", CELL)]]
    for g, lns in groups.items():
        nos = ", ".join(str(l["no"]) for l in lns)
        rows.append([Paragraph(g, CELL), Paragraph(nos, CELL),
                     Paragraph(str(sum(l["qty"] for l in lns)), CELL),
                     Paragraph(", ".join(l["sku"] for l in lns), CELL)])
    t = Table(rows, colWidths=[32 * mm, 40 * mm, 16 * mm, 82 * mm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2F4F6F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#AAAAAA")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    fl += [t, Spacer(1, 8)]

    fl.append(Paragraph("<b>3. Delivery locations and split</b>", BODY))
    fl.append(Paragraph(
        "Enduser compute delivers to all three sites in the proportion 60:25:15 "
        "(Bengaluru : Pune : Gurugram). Datacentre and network hardware delivers to "
        "Pune (DC-West) and Gurugram (DC-North) only, split evenly. Retail and "
        "peripheral items deliver to Pune and Gurugram distribution centres. Exact "
        "site-wise quantities issue with the purchase order.", BODY))
    fl.append(Spacer(1, 7))

    fl.append(Paragraph("<b>4. Included in scope</b>", BODY))
    for t_ in ["Delivery to each site, unloading and placement.",
               "Unboxing, asset tagging against our asset register format, and "
               "recording of serial numbers in a returnable sheet.",
               "Racking, stacking, cabling and power-up of datacentre and network "
               "hardware. Firmware brought to the vendor's current recommended level.",
               "Windows imaging of enduser devices from an image we supply.",
               "Warranty and onsite support for the term stated in each line "
               "description, from the date of acceptance.",
               "Collection and certified disposal of packaging."]:
        fl.append(Paragraph(f"• {t_}", BODY))
    fl.append(Spacer(1, 7))

    fl.append(Paragraph("<b>5. Excluded from scope</b>", BODY))
    for t_ in ["Structured cabling and containment beyond patching to existing ports.",
               "Electrical work upstream of the rack PDU.",
               "Application software, licences and migration of user data.",
               "Buyback of the decommissioned fleet, which is being tendered "
               "separately. Question 10 of the questionnaire asks about capability "
               "only and is not priced here."]:
        fl.append(Paragraph(f"• {t_}", BODY))
    fl.append(Spacer(1, 7))

    fl.append(Paragraph("<b>6. Indicative timeline</b>", BODY))
    rows = [[Paragraph(f"<b>{h}</b>", CELL) for h in ["Milestone", "Date"]]]
    for m, dt in [("Enquiry issued", f"{RFX['issued']:%d %b %Y}"),
                  ("Clarification questions close", "11 Sep 2026"),
                  ("Responses due", f"{RFX['due']:%d %b %Y}, 17:00 IST"),
                  ("Commercial evaluation and clarifications", "19 to 24 Sep 2026"),
                  ("Award recommendation to VP Procurement", "26 Sep 2026"),
                  ("Purchase order", "30 Sep 2026"),
                  ("Delivery window opens", "15 Oct 2026"),
                  ("Full deployment complete", "31 Jan 2027")]:
        rows.append([Paragraph(m, CELL), Paragraph(dt, CELL)])
    t = Table(rows, colWidths=[110 * mm, 60 * mm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2F4F6F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#AAAAAA")),
    ]))
    fl += [t, Spacer(1, 8)]

    fl.append(Paragraph("<b>7. Award basis</b>", BODY))
    fl.append(Paragraph(
        "We reserve the right to award by line, by group, or in whole, to one or "
        "more vendors. Lowest price is not the sole criterion. A bid that fails any "
        "mandatory question in <i>03_Vendor_Questionnaire.xlsx</i> will not be "
        "considered at any price.", BODY))

    fl += _sign_off()
    _doc(p, "Scope of Work").build(fl, onFirstPage=_stamp("Scope of Work"),
                                   onLaterPages=_stamp("Scope of Work"))
    written.append(p)


# ===========================================================================
# 02  Line Items
# ===========================================================================

def line_items():
    p = OUT / "02_Line_Items.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Line Items"

    ws["A1"] = f"{RFX['id']} - Line Items"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"{BUYER['legal_name']}  |  GSTIN {BUYER['gstin']}"
    ws["A3"] = (f"Issued {RFX['issued']:%d %b %Y}   |   Responses due "
                f"{RFX['due']:%d %b %Y}, 17:00 IST")
    ws["A4"] = ("Quote ex-GST in INR. State the GST rate and HSN against every line. "
                f"{RFX['incoterm_asked']}.")
    ws["A5"] = ("DO NOT ALTER THE UNIT OF MEASURE. Lines shown in red are priced per "
                "kit, pack or box. Quote per that unit, not per piece.")
    ws["A5"].font = Font(bold=True, color=RED)
    ws["A6"] = ("Mark any line you cannot supply as NQ. Do not leave it blank: a blank "
                "line and a declined line are read differently.")
    ws["A6"].font = Font(bold=True, color=RED)

    hdr = ["Line", "SKU", "Group", "Description", "Key spec", "UoM",
           "Units per UoM", "Qty", "HSN", "Unit rate (INR, ex-GST)",
           "Line total (INR)", "GST %", "Lead time (weeks)",
           "Make / model offered", "Remarks"]
    r = 8
    for c, h in enumerate(hdr, 1):
        cell = ws.cell(r, c, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BOX

    for i, ln in enumerate(LINES):
        rr = r + 1 + i
        spec = ", ".join(f"{k}={v}" for k, v in ln["spec"].items())
        for c, val in enumerate([ln["no"], ln["sku"], ln["group"], ln["desc"], spec,
                                 ln["uom"], ln["pack_size"], ln["qty"], ln["hsn"],
                                 None, None, None, None, None, None], 1):
            cell = ws.cell(rr, c, val)
            cell.border = BOX
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(rr, 11).value = f'=IF(J{rr}="","",H{rr}*J{rr})'
        if ln["pack_size"] > 1:
            for c in (6, 7):
                ws.cell(rr, c).font = Font(bold=True, color=RED)
            ws.cell(rr, 6).fill = PatternFill("solid", fgColor="FDE9E9")
            ws.cell(rr, 7).fill = PatternFill("solid", fgColor="FDE9E9")

    tr = r + 1 + len(LINES)
    ws.cell(tr, 9, "TOTAL (ex-GST)").font = Font(bold=True)
    ws.cell(tr, 11, f"=SUM(K{r+1}:K{tr-1})").font = Font(bold=True)
    ws.cell(tr + 2, 1, "Vendor name:").font = Font(bold=True)
    ws.cell(tr + 3, 1, "Signed by / date:").font = Font(bold=True)
    ws.cell(tr + 4, 1, "Bid validity offered:").font = Font(bold=True)

    for c, w in enumerate([6, 12, 20, 52, 40, 12, 8, 7, 8, 15, 15, 7, 13, 30, 30], 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A9"

    notes = wb.create_sheet("Read this first")
    notes["A1"] = "Four things that cause most of the rework"
    notes["A1"].font = Font(bold=True, size=13)
    for i, t_ in enumerate([
        "1. Unit of measure. Line 13 is one KIT of 2x 16GB DIMMs, not one DIMM. "
        "Line 17 is one PACK of 5 cartridges. Line 20 is one BOX of 10 "
        "transceivers. Line 22 is one BOX of 50 patch cords. Quote per that unit.",
        "2. Warranty. Where a line description states a warranty term, price it "
        "inclusive of that term. Do not quote a shorter term and add an uplift on "
        "a separate line.",
        "3. Equivalents. Permitted, but state the offered make and model in the "
        "'Make / model offered' column. An unstated substitution is treated as "
        "non-compliant.",
        "4. Conditional discounts. Any discount that depends on an approval, an "
        "order date or a competitor's price must be labelled as conditional. "
        "Conditional discounts are recorded but are not used in ranking.",
    ]):
        c = notes.cell(3 + i * 2, 1, t_)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        notes.merge_cells(start_row=3 + i * 2, start_column=1,
                          end_row=3 + i * 2, end_column=8)
        notes.row_dimensions[3 + i * 2].height = 44
    notes.column_dimensions["A"].width = 108

    wb.save(p)
    written.append(p)


# ===========================================================================
# 03  Vendor Questionnaire
# ===========================================================================

def questionnaire():
    p = OUT / "03_Vendor_Questionnaire.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Questionnaire"

    ws["A1"] = f"{RFX['id']} - Vendor Questionnaire"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"{BUYER['legal_name']}  |  Return with your commercial offer."
    ws["A3"] = ("M = MANDATORY. Failure OR non-response on ANY mandatory question "
                "disqualifies the bid regardless of price. D = desirable, scored.")
    ws["A3"].font = Font(bold=True, color=RED)
    ws["A4"] = ("Where a document is required, attach it and name the file in the "
                "last column. An answer that its own attachment contradicts will be "
                "read against the attachment.")
    ws["A4"].font = Font(bold=True, color=RED)

    hdr = ["Q", "M/D", "Question", "Your answer", "Document required",
           "Attached filename"]
    for c, h in enumerate(hdr, 1):
        cell = ws.cell(6, c, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.border = BOX
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for i, q in enumerate(QUESTIONNAIRE):
        rr = 7 + i
        for c, val in enumerate([q["no"], q["kind"], q["q"], None,
                                 "Yes" if q["doc_required"] else "No", None], 1):
            cell = ws.cell(rr, c, val)
            cell.border = BOX
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        if q["kind"] == "M":
            ws.cell(rr, 2).font = Font(bold=True, color=RED)
            ws.cell(rr, 2).fill = PatternFill("solid", fgColor="FDE9E9")
        ws.row_dimensions[rr].height = 34

    n_m = sum(1 for q in QUESTIONNAIRE if q["kind"] == "M")
    end = 7 + len(QUESTIONNAIRE)
    ws.cell(end + 1, 1,
            f"{len(QUESTIONNAIRE)} questions. {n_m} mandatory, "
            f"{len(QUESTIONNAIRE) - n_m} desirable.").font = Font(bold=True)
    ws.cell(end + 3, 1, "Authorised signatory:").font = Font(bold=True)
    ws.cell(end + 4, 1, "Name and designation:").font = Font(bold=True)
    ws.cell(end + 5, 1, "Date:").font = Font(bold=True)

    for c, w in enumerate([7, 7, 92, 46, 18, 34], 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A7"

    wb.save(p)
    written.append(p)


# ===========================================================================
# 04  Commercial Terms
# ===========================================================================

def commercial_terms():
    p = OUT / "04_Commercial_Terms.pdf"
    fl = _letterhead("Commercial Terms")
    fl.append(Paragraph(
        "These terms apply to any purchase order placed against this enquiry. "
        "Counter-terms may be offered and will be evaluated as a commercial "
        "deviation. Silence is acceptance.", BODY))
    fl.append(Spacer(1, 8))

    rows = [[Paragraph("<b>Item</b>", CELL), Paragraph("<b>Our requirement</b>", CELL)]]
    for k, v in [
        ("Currency", f"{RFX['currency_asked']} only. A quotation in any other "
                     f"currency must state the reference rate used, its source and "
                     f"its date, and must say whether the rate is committed."),
        ("Tax basis", RFX["tax_basis_asked"] + " Quote GST separately. Do not quote "
                      "a tax-inclusive rate."),
        ("Delivery basis", f"{RFX['incoterm_asked']}. Freight, insurance and "
                           f"unloading to the vendor's account. Ex-works quotations "
                           f"will be evaluated but the buyer will not estimate duty "
                           f"or freight on the vendor's behalf."),
        ("Payment", f"{RFX['payment_asked']}. Advance payment is not offered. "
                    f"Milestone payment may be discussed for the datacentre group."),
        ("Bid validity", f"{RFX['validity_asked']}. A validity shorter than 30 days "
                         f"is likely to expire before our approval cycle closes."),
        ("Price firmness", "Prices to be firm and fixed for the full delivery "
                           "period. Price variation, forex variation and "
                           "'subject to availability' clauses are commercial "
                           "deviations and will be recorded as such."),
        ("Warranty", "As stated in each line description, from the date of "
                     "acceptance and not the date of despatch. Onsite, at our "
                     "premises, with no per-incident charge."),
        ("Support SLA", "Next business day onsite for enduser hardware. Four hours "
                        "onsite for datacentre and network hardware. Both at every "
                        "location listed in your answer to Question 5."),
        ("Acceptance", "Within ten working days of delivery, against serial-number "
                       "verification and successful power-up. Payment terms run "
                       "from acceptance."),
        ("Liquidated damages", "0.5% of the value of the delayed portion per "
                               "completed week, capped at 5%. Not applicable where "
                               "the delay is attributable to us."),
        ("Partial award", "We may award by line, by group or in whole, to one or "
                          "more vendors. A quotation valid only on an all-or-nothing "
                          "basis must say so explicitly."),
        ("Conditional discounts", "Any discount conditional on an approval, an order "
                                  "date or another bidder's price must be labelled "
                                  "conditional. Such discounts are recorded and "
                                  "shown to the approver, but are not used in the "
                                  "ranked comparison."),
        ("Minimum order qty", "State any MOQ against the line. Where MOQ exceeds our "
                              "quantity, we will evaluate your unit rate against OUR "
                              "quantity and treat the excess as a separate decision."),
        ("Free goods", "Bundled or free-of-cost items must be listed with a stated "
                       "value. A free item with no stated value cannot be compared."),
        ("Compliance", "E-waste handling under the E-Waste (Management) Rules 2022. "
                       "BIS registration where applicable. Original, unused, "
                       "India-market stock with full manufacturer warranty."),
        ("Sub-contracting", "Not permitted for warranty and support without our "
                            "prior written consent."),
        ("Confidentiality", "This enquiry and its contents are confidential. Do not "
                            "disclose our requirement, quantities or your pricing to "
                            "any third party including other bidders."),
        ("Governing law", "Laws of India. Courts at Bengaluru have exclusive "
                          "jurisdiction."),
    ]:
        rows.append([Paragraph(f"<b>{k}</b>", CELL), Paragraph(v, CELL)])

    t = Table(rows, colWidths=[38 * mm, 132 * mm], repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2F4F6F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#AAAAAA")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1),
         [colors.white, colors.HexColor("#F4F6F8")]),
    ]))
    fl += [t]
    fl += _sign_off()
    _doc(p, "Commercial Terms").build(fl, onFirstPage=_stamp("Commercial Terms"),
                                      onLaterPages=_stamp("Commercial Terms"))
    written.append(p)


# ===========================================================================
# 05  Instructions to Bidders
# ===========================================================================

def instructions():
    p = OUT / "05_Instructions_to_Bidders.pdf"
    c = RFX["buyer_contact"]
    fl = _letterhead("Instructions to Bidders")
    fl.append(Paragraph(
        "This enquiry comprises five documents: the Scope of Work, the Line Items "
        "sheet, the Vendor Questionnaire, the Commercial Terms, and these "
        "instructions. Read all five before quoting.", BODY))
    fl.append(Spacer(1, 8))

    for i, t_ in enumerate([
        f"Responses are due by {RFX['due']:%d %B %Y} at 17:00 IST, by email to "
        f"{c['email']}. Clarification questions close 11 September 2026.",
        "Return the Line Items sheet and the Questionnaire as separate files. "
        "Quotations in your own format will be accepted, but the more your format "
        "differs from ours, the longer evaluation takes and the more likely we are "
        "to come back to you with clarifications.",
        f"Quote in {RFX['currency_asked']}, ex-GST, {RFX['incoterm_asked']}. State "
        f"the GST rate and HSN code against every line.",
        "<b>Do not alter the unit of measure.</b> Line 13 is one kit of 2x 16GB "
        "DIMMs. Line 17 is one pack of 5 cartridges. Line 20 is one box of 10 "
        "transceivers. Line 22 is one box of 50 patch cords. A price per piece "
        "against a line specified per box will be normalised, and you will be asked "
        "to confirm, which costs us both time.",
        "Partial bids are permitted. Mark any line you cannot supply as <b>NQ</b>. "
        "Do not leave it blank. A blank line and a declined line are treated "
        "differently: the first makes us chase you, the second does not.",
        "Equivalents are permitted provided the offered make and model is stated in "
        "the 'Make / model offered' column and meets or exceeds the key spec. An "
        "unstated substitution will be treated as non-compliant.",
        "Where a line description states a warranty term, quote inclusive of that "
        "term. Do not quote a shorter term with a separate uplift line: it makes "
        "your bid look cheaper than it is and it will be adjusted onto a common "
        "basis before comparison.",
        "Discounts applied only at total level cannot be attributed to individual "
        "lines and therefore cannot be used where we award by line. If you want a "
        "discount to count on a partial award, apply it to the line rates.",
        "Any discount conditional on an approval you do not yet hold, an order date, "
        "or another bidder's price must be stated as conditional. It will be recorded "
        "and shown to our approver, but it will not be used in the ranked comparison.",
        "<b>All mandatory questionnaire questions must be answered and evidenced. "
        "Failure or non-response on any mandatory question disqualifies the bid "
        "regardless of price.</b> Where an answer and its attachment disagree, the "
        "attachment governs. Check the validity dates on your certificates before "
        "you attach them.",
        f"Revised quotations are accepted until the due date. Mark them clearly as a "
        f"revision, state the quotation number they supersede, and identify which "
        f"lines changed. An unmarked revision may be evaluated alongside your "
        f"original rather than instead of it.",
        "Do not price by reference to a previous purchase order, a percentage uplift "
        "on last year, or a promise to match another bidder. We cannot evaluate a "
        "price we have to derive, and a price contingent on a competitor's bid "
        "cannot be ranked or awarded.",
    ], 1):
        fl.append(Paragraph(f"{i}. {t_}", BODY))
        fl.append(Spacer(1, 5))

    fl += _sign_off()
    _doc(p, "Instructions to Bidders").build(
        fl, onFirstPage=_stamp("Instructions to Bidders"),
        onLaterPages=_stamp("Instructions to Bidders"))
    written.append(p)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    # Remove the old combined file if a previous run left it behind.
    for stale in ("RFX-2026-0142_Line_Items_and_Questionnaire.xlsx",
                  "RFX-2026-0142_Instructions_to_Bidders.pdf"):
        q = OUT / stale
        if q.exists():
            q.unlink()
            print(f"  removed superseded {stale}")

    scope_of_work()
    line_items()
    questionnaire()
    commercial_terms()
    instructions()

    print(f"\nRFx pack: {len(written)} documents in {OUT.relative_to(OUT.parents[2])}/")
    for q in sorted(written):
        print(f"  {q.stat().st_size/1024:7.1f} KB  {q.name}")


if __name__ == "__main__":
    main()
