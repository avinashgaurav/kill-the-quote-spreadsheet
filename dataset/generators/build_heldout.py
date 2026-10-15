"""
A held-out corpus: five new vendor replies the reader prompt has never seen.

The original five documents were also the documents the prompt was tuned on, so
a perfect score there is a score on the training set. This builds a second set
against the same RFx (the buyer's enquiry is an input, not an answer) with
different suppliers, prices, layouts and line selection, and writes the answer
key from the same table the documents are rendered from.

    python3 build_heldout.py

Writes dataset/heldout/ (the documents, answer-key.json, manifest.json).
Everything is fictional: suppliers, people, addresses, GSTINs and prices.

  H1  Kestrel Compute      xlsx  two sheets, own item codes, GST column, one
                                 regret, one substitution, one per-piece line
  H2  Tarang Datacom       pdf   two pages, rate and amount columns side by side,
                                 one USD line, one per-cartridge line, one regret
  H3  Nilgiri Tech Traders docx  a letter, Indian digit grouping in prose, a
                                 per-pair line, one complimentary (zero) line
  H4  Tejas Systems        jpg   printed quotation sheet photographed twice
                                 (moderate, very hard), two pen overrides, one
                                 price under glare
  H5  Saptarishi Infra     eml   plain-text reply, a per-box-of-25 line, a
                                 relative price, a regret
"""

import json
from datetime import datetime
from email.message import EmailMessage
from email.utils import format_datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from docx import Document
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                TableStyle, PageBreak)
from PIL import Image, ImageDraw, ImageFont

from catalog import LINES, RFX, BUYER
from build_photo import glare, perspective
from build_photo_variants import shoot, curl

OUT = Path(__file__).resolve().parents[1] / "heldout"
LINE = {ln["no"]: ln for ln in LINES}


def q(price, uom="nos", status="quoted", ccy="INR", **kw):
    return dict(status=status, price=price, uom=uom, ccy=ccy, **kw)


# ---------------------------------------------------------------------------
# The answer key. Every document below is rendered from this table.
# ---------------------------------------------------------------------------

KEY = {
    "H1": {
        1: q(63400), 2: q(108900), 4: q(40650), 5: q(9620), 6: q(21050),
        8: q(1415), 9: q(1795), 10: q(3150),
        12: q(562000, status="substituted",
              offered="Dell PowerEdge R760, 2x Xeon Silver 4410Y, 128GB"),
        13: q(18350, "kit"),
        14: dict(status="not_quoted", note="Regret, not available from our OEM."),
        18: q(189500), 20: q(40800, "box of 10"), 21: q(28150),
        22: q(262, "per pc", status="uom_mismatch",
              note="Priced per patch cord against a box of 50 asked."),
        23: q(77300), 25: q(66900), 27: q(9280), 28: q(51750),
    },
    "H2": {
        12: q(574500), 13: q(18990, "kit"), 14: q(66800), 15: q(829000),
        16: dict(status="not_quoted", note="Not offered."),
        17: q(4480, "per cartridge", status="uom_mismatch",
              note="Priced per cartridge against a pack of 5 asked."),
        18: q(193200),
        19: q(5240, ccy="USD", note="USD, ex-works Singapore."),
        20: q(42600, "box of 10"), 21: q(27650), 22: q(14050, "box of 50"),
        23: q(79900), 24: q(306500), 25: q(67250),
    },
    "H3": {
        1: q(61750), 3: q(88400), 5: q(9380), 7: q(13250), 9: q(1720),
        10: q(6180, "per pair", status="uom_mismatch",
              note="Headsets priced per pair against single units asked."),
        11: q(2640), 26: q(146500), 27: q(9050), 29: q(46300),
        30: q(0, note="Years 4 and 5 cover complimentary on laptops. Zero-price line."),
    },
    "H4": {
        1: q(62300), 2: q(106400), 3: q(87900), 5: q(9540), 6: q(20750),
        7: q(13050), 8: q(1360),
        9: q(1640, handwritten=True, printed_price=1690,
             note="Printed 1,690 struck through, 1,640 written in pen."),
        11: q(2580), 13: q(18700, "kit"), 18: q(186500),
        20: q(41900, "box of 10"),
        21: dict(status="illegible", printed_price=27800,
                 note="Glare across the cell. A price is present and cannot be read."),
        27: q(9120), 28: q(50900),
        29: q(45600, handwritten=True, printed_price=46900,
              note="Printed 46,900 struck through, 45,600 written in pen."),
    },
    "H5": {
        4: q(39950), 6: q(20600), 8: q(1290),
        21: dict(status="relative", basis="prior_invoice",
                 note="'access points at the same rate as our last invoice'"),
        22: q(7250, "per box of 25", status="uom_mismatch",
              note="Priced per box of 25 against a box of 50 asked."),
        23: dict(status="not_quoted", note="'racks we cannot supply this time'"),
        25: q(65400),
    },
}

VENDORS = {
    "H1": dict(name="Kestrel Compute Pvt Ltd", city="Coimbatore",
               contact="Divya Raghunathan", email="divya.r@kestrelcompute.example",
               gstin="33AAKCK4410P1Z7"),
    "H2": dict(name="Tarang Datacom Pvt Ltd", city="Bhubaneswar",
               contact="Subhrajit Mohanty", email="s.mohanty@tarangdatacom.example",
               gstin="21AADCT7782M1Z2"),
    "H3": dict(name="Nilgiri Tech Traders", city="Udhagamandalam",
               contact="Joseph Mathai", email="joseph@nilgiritech.example",
               gstin="33AAQFN5530R1Z9"),
    "H4": dict(name="Tejas Systems", city="Madurai",
               contact="M. Senthilnathan", email="tejas.systems@mail.example",
               gstin="33AAHFT2291K1ZD"),
    "H5": dict(name="Saptarishi Infra Services", city="Nagpur",
               contact="Ketan Wankhede", email="ketan@saptarishi-infra.example",
               gstin="27AAKFS6614B1Z5"),
}

FILES = {
    "H1": "H1_Kestrel_PriceSchedule_KC-2026-311.xlsx",
    "H2": "H2_Tarang_Offer_TDC-OF-0926-77.pdf",
    "H3": "H3_Nilgiri_Quotation_Letter.docx",
    "H4a": "H4_tejas_quote_photo_moderate.jpg",
    "H4b": "H4_tejas_quote_photo_lowlight.jpg",
    "H5": "H5_saptarishi_reply.eml",
}


def inr(n):
    """Indian digit grouping: 146500 -> 1,46,500."""
    s = str(int(n))
    if len(s) <= 3:
        return s
    head, tail = s[:-3], s[-3:]
    parts = []
    while len(head) > 2:
        parts.insert(0, head[-2:])
        head = head[:-2]
    if head:
        parts.insert(0, head)
    return ",".join(parts + [tail])


def priced(code):
    return {n: r for n, r in KEY[code].items() if r.get("price") is not None}


# ---------------------------------------------------------------------------
# H1: xlsx
# ---------------------------------------------------------------------------

H1_CODES = {1: "KC-NB-1405", 2: "KC-NB-1471", 4: "KC-DT-0512", 5: "KC-MN-2401",
            6: "KC-MN-2702", 8: "KC-AC-BP15", 9: "KC-AC-KM01", 10: "KC-AC-HS20",
            12: "KC-SV-R760", 13: "KC-SV-MK32", 14: "KC-SV-SD19", 18: "KC-NW-S48P",
            20: "KC-NW-SR10", 21: "KC-NW-AP6X", 22: "KC-NW-PC6A", 23: "KC-PW-R42U",
            25: "KC-PW-U3K", 27: "KC-RT-SC2D", 28: "KC-RT-HH13"}

H1_GROUPS = [("CLIENT DEVICES", [1, 2, 4, 5, 6, 8, 9, 10]),
             ("SERVER & STORAGE", [12, 13, 14]),
             ("NETWORK, RACK & POWER", [18, 20, 21, 22, 23, 25]),
             ("STORE OPERATIONS", [27, 28])]


def build_h1():
    v = VENDORS["H1"]
    wb = openpyxl.Workbook()
    cover = wb.active
    cover.title = "Cover"
    for i, line in enumerate([
        v["name"], f"{v['city']}  |  GSTIN {v['gstin']}", "",
        f"Price schedule KC-2026-311 against {RFX['id']}",
        f"Attention: {RFX['buyer_contact']['name']}, {BUYER['short_name']}",
        "Date: 12-Sep-2026", "",
        "Commercial terms",
        "a) Rates are ex-GST; GST extra as shown per line.",
        "b) Delivery FOR all three sites, 4 to 6 weeks.",
        "c) Payment 30 days from delivery.",
        "d) Validity 60 days.",
        "e) Price schedule is on the next sheet.",
    ], start=1):
        cover.cell(i, 1, line)
    cover["A1"].font = Font(bold=True, size=14)

    ws = wb.create_sheet("Price Schedule")
    hdr = ["Our Code", "RFx Ln", "Item", "Make / Model", "Unit", "Qty",
           "Rate ex-GST (INR)", "GST %", "Line Value (INR)"]
    for c, h in enumerate(hdr, start=1):
        cell = ws.cell(2, c, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F4F4F")
        cell.alignment = Alignment(wrap_text=True)
    ws.cell(1, 1, "KESTREL COMPUTE: PRICE SCHEDULE KC-2026-311").font = Font(bold=True)

    r = 3
    total = 0
    for group, nos in H1_GROUPS:
        ws.cell(r, 1, group).font = Font(bold=True)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
        r += 1
        for n in nos:
            t = KEY["H1"][n]
            ln = LINE[n]
            ws.cell(r, 1, H1_CODES[n])
            ws.cell(r, 2, n)
            ws.cell(r, 3, ln["desc"].split(",")[0])
            ws.cell(r, 4, t.get("offered", "As specified"))
            if t["status"] == "not_quoted":
                ws.cell(r, 5, "-")
                ws.cell(r, 6, ln["qty"])
                ws.cell(r, 7, "Regret")
                ws.cell(r, 9, "Not available from our OEM")
            else:
                unit = {"nos": "Each", "kit": "Kit (2x16GB)", "box of 10": "Box/10",
                        "per pc": "Per piece"}[t["uom"]]
                qty = ln["qty"] * (50 if t["uom"] == "per pc" else 1)
                ws.cell(r, 5, unit)
                ws.cell(r, 6, qty)
                ws.cell(r, 7, t["price"]).number_format = "#,##0"
                ws.cell(r, 8, 18 if n not in (8,) else 12)
                ws.cell(r, 9, t["price"] * qty).number_format = "#,##0"
                total += t["price"] * qty
            r += 1
    r += 1
    ws.cell(r, 8, "Total ex-GST").font = Font(bold=True)
    ws.cell(r, 9, total).number_format = "#,##0"
    r += 2
    ws.cell(r, 1, "Notes: Line 12 offered on Dell PowerEdge R760 in place of the "
                  "referenced platform. Patch cords are priced per piece; "
                  "quantity shown is 14 boxes x 50.")
    for col, w in zip("ABCDEFGHI", [13, 7, 40, 34, 13, 7, 16, 7, 17]):
        ws.column_dimensions[col].width = w
    wb.save(OUT / FILES["H1"])


# ---------------------------------------------------------------------------
# H2: pdf
# ---------------------------------------------------------------------------

def build_h2():
    v = VENDORS["H2"]
    ss = getSampleStyleSheet()
    small = ParagraphStyle("small", parent=ss["Normal"], fontSize=8, leading=10)
    body = ParagraphStyle("body", parent=ss["Normal"], fontSize=9.5, leading=13)
    doc = SimpleDocTemplate(str(OUT / FILES["H2"]), pagesize=A4,
                            leftMargin=16 * mm, rightMargin=16 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm)
    story = [
        Paragraph(f"<b>{v['name'].upper()}</b>", ss["Title"]),
        Paragraph(f"Plot 17, Chandaka Industrial Estate, {v['city']} 751024  |  "
                  f"GSTIN {v['gstin']}", small),
        Spacer(1, 6 * mm),
        Paragraph("<b>Techno-commercial offer TDC/OF/0926/77</b>  dated 14 September 2026",
                  body),
        Paragraph(f"To: {RFX['buyer_contact']['name']}, {BUYER['legal_name']}. "
                  f"Ref: {RFX['id']}.", body),
        Spacer(1, 3 * mm),
        Paragraph("We are pleased to offer the datacentre and network items below. "
                  "Unit rates are ex-GST. The Amount column is unit rate times the "
                  "quantity in your enquiry and is given for reference only.", body),
        Spacer(1, 4 * mm),
    ]

    def section(title, nos):
        rows = [["Sl", "Your line", "Description (abridged)", "Qty", "UoM",
                 "Unit rate", "Amount"]]
        for i, n in enumerate(nos, start=1):
            t = KEY["H2"][n]
            ln = LINE[n]
            desc = ln["desc"].split(",")[0]
            if t["status"] == "not_quoted":
                rows.append([str(i), f"L{n}", desc, str(ln["qty"]), "-",
                             "Not offered", "-"])
                continue
            uom = {"nos": "No.", "kit": "Kit", "box of 10": "Box of 10",
                   "box of 50": "Box of 50", "per cartridge": "Cartridge"}[t["uom"]]
            qty = ln["qty"] * (5 if t["uom"] == "per cartridge" else 1)
            cur = "USD " if t["ccy"] == "USD" else "Rs "
            rate = f"{t['price']:,}" if t["ccy"] == "USD" else inr(t["price"])
            amt = (f"{t['price'] * qty:,}" if t["ccy"] == "USD"
                   else inr(t["price"] * qty))
            rows.append([str(i), f"L{n}", desc + (" *" if t["ccy"] == "USD" else ""),
                         str(qty), uom, cur + rate, cur + amt])
        tbl = Table(rows, colWidths=[9 * mm, 15 * mm, 66 * mm, 12 * mm, 20 * mm,
                                     23 * mm, 30 * mm], repeatRows=1)
        tbl.setStyle(TableStyle([
            ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8.5),
            ("FONT", (0, 1), (-1, -1), "Helvetica", 8.5),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDE6EE")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#8899AA")),
            ("ALIGN", (5, 1), (-1, -1), "RIGHT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        return [Paragraph(f"<b>{title}</b>", body), Spacer(1, 2 * mm), tbl,
                Spacer(1, 5 * mm)]

    story += section("Section A: Network, rack and power", [18, 19, 20, 21, 22, 23, 24, 25])
    story += [PageBreak()]
    story += section("Section B: Servers and storage", [12, 13, 14, 15, 16, 17])
    story += [
        Paragraph("* Line 19 is quoted in US dollars, ex-works Singapore. Customs "
                  "duty, freight and insurance to your account.", small),
        Spacer(1, 4 * mm),
        Paragraph("<b>Terms.</b> Payment 45 days from invoice. Delivery 6 to 8 weeks. "
                  "Validity 45 days. LTO-9 media is sold per cartridge only.", body),
        Spacer(1, 10 * mm),
        Paragraph(f"For {v['name']}<br/><br/>{v['contact']}<br/>Head, Enterprise Sales",
                  body),
    ]
    doc.build(story)


# ---------------------------------------------------------------------------
# H3: docx letter
# ---------------------------------------------------------------------------

def build_h3():
    v = VENDORS["H3"]
    k = KEY["H3"]
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Georgia"
    st.font.size = Pt(10.5)
    p = d.add_paragraph()
    p.add_run(v["name"].upper()).bold = True
    d.add_paragraph(f"Commercial Road, {v['city']} 643001  |  GSTIN {v['gstin']}")
    d.add_paragraph("16 September 2026")
    d.add_paragraph(f"{RFX['buyer_contact']['name']}\n{RFX['buyer_contact']['title']}\n"
                    f"{BUYER['legal_name']}")
    d.add_paragraph(f"Sub: Our quotation against your enquiry {RFX['id']}")
    d.add_paragraph("Dear Madam,")
    d.add_paragraph(
        "Thank you for including us. We have gone through the line items and are "
        "happy to quote for the portion of the requirement we handle regularly. All "
        "rates below are per unit, exclusive of GST, delivered to your three sites.")
    d.add_paragraph(
        f"For laptops, we can supply the 14-inch business model (your line 1) at "
        f"Rs. {inr(k[1]['price'])}/- each and the 13.3-inch ultraportable (line 3) at "
        f"Rs. {inr(k[3]['price'])}/- each, both with three years onsite warranty. "
        f"Further, cover for years four and five on these laptops (your line 30) "
        f"will be provided complimentary, at no additional charge.")
    d.add_paragraph(
        f"On displays and docks, the 24-inch monitor (line 5) is Rs. "
        f"{inr(k[5]['price'])}/- and the USB-C docking station (line 7) is Rs. "
        f"{inr(k[7]['price'])}/-. The keyboard and mouse combo (line 9) is "
        f"Rs. {inr(k[9]['price'])}/- and the 1080p webcam (line 11) Rs. "
        f"{inr(k[11]['price'])}/-.")
    d.add_paragraph(
        f"Please note that the noise-cancelling headsets (line 10) are packed and "
        f"sold by the pair; our rate is Rs. {inr(k[10]['price'])}/- per pair.")
    d.add_paragraph("For the store and print items we attach a short table.")
    t = d.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, ["Your line", "Item", "Rate per unit (Rs.)"]):
        c.text = h
    for n in (26, 27, 29):
        row = t.add_row().cells
        row[0].text = str(n)
        row[1].text = LINE[n]["desc"].split(",")[0]
        row[2].text = inr(k[n]["price"])
    d.add_paragraph("")
    d.add_paragraph(
        "We have not quoted for servers, storage, network or power items. Payment "
        "terms: 50% advance, balance on delivery. Offer valid for 30 days.")
    d.add_paragraph("Yours faithfully,")
    d.add_paragraph(f"{v['contact']}\nProprietor")
    d.save(OUT / FILES["H3"])


# ---------------------------------------------------------------------------
# H4: photographed quotation sheet
# ---------------------------------------------------------------------------

F = "/System/Library/Fonts/Supplemental/"
H4_SHORT = {1: "Laptop 14 i5 16G 512G", 2: "Laptop 14 i7 32G 1T", 3: "Laptop 13.3 ultra i5",
            5: "Monitor 24 FHD", 6: "Monitor 27 QHD USB-C", 7: "USB-C dock 100W",
            8: "Backpack 15.6", 9: "Wireless KB/Mouse", 11: "Webcam 1080p",
            13: "DDR5 RDIMM 2x16G kit", 18: "48P PoE+ switch", 20: "10G SR SFP+",
            21: "WiFi 6 AP w/ 3y lic", 27: "2D BT scanner+cradle",
            28: "Rugged HHT Android", 29: "55in signage 24x7"}
H4_UNIT = {"nos": "Nos", "kit": "Kit", "box of 10": "Box/10"}
PAGE_W, PAGE_H = 1700, 1900
INK, PEN, PAPER = (30, 30, 34), (20, 40, 125), (247, 246, 240)


def render_h4():
    f = lambda name, size: ImageFont.truetype(F + name, size)
    img = Image.new("RGB", (PAGE_W, PAGE_H), PAPER)
    d = ImageDraw.Draw(img)
    v = VENDORS["H4"]
    d.text((90, 70), "TEJAS SYSTEMS", font=f("Verdana Bold.ttf", 56), fill=(20, 90, 60))
    d.text((92, 146), f"Computers & Networking  |  {v['city']}  |  GSTIN {v['gstin']}",
           font=f("Verdana.ttf", 24), fill=(90, 90, 95))
    d.line([90, 196, PAGE_W - 90, 196], fill=(20, 90, 60), width=4)
    d.text((90, 230), "QUOTATION", font=f("Trebuchet MS Bold.ttf", 40), fill=INK)
    d.text((90, 286), f"To: {BUYER['short_name']}   Enq: {RFX['id']}   Dt: 15/09/2026",
           font=f("Verdana.ttf", 25), fill=INK)

    x0, y = 90, 360
    col = [x0, x0 + 110, x0 + 760, x0 + 960, PAGE_W - 90]
    d.rectangle([x0, y, PAGE_W - 90, y + 54], fill=(225, 235, 228))
    hf = f("Verdana Bold.ttf", 24)
    for cx, label in zip(col, ["LINE", "ITEM", "UNIT", "RATE Rs."]):
        d.text((cx + 14, y + 14), label, font=hf, fill=INK)
    y += 54
    bf, mf = f("Verdana.ttf", 25), f("Andale Mono.ttf", 27)
    rows = []
    for n in sorted(KEY["H4"]):
        t = KEY["H4"][n]
        h = 54
        d.line([x0, y + h, PAGE_W - 90, y + h], fill=(205, 210, 207), width=2)
        d.text((col[0] + 24, y + 14), str(n), font=bf, fill=INK)
        d.text((col[1] + 14, y + 14), H4_SHORT[n], font=bf, fill=INK)
        d.text((col[2] + 14, y + 14), H4_UNIT[t.get("uom", "nos")], font=bf, fill=INK)
        txt = inr(t.get("printed_price", t.get("price")))
        w = d.textlength(txt, font=mf)
        d.text((col[4] - 24 - w, y + 13), txt, font=mf, fill=INK)
        rows.append((n, y, y + h, col[4] - 24 - w, w))
        y += h
    d.rectangle([x0, 360, PAGE_W - 90, y], outline=(130, 140, 135), width=3)
    for cx in col[1:-1]:
        d.line([cx, 360, cx, y], fill=(190, 195, 192), width=2)
    fy = y + 40
    for i, line in enumerate(["Rates per unit, GST extra @ 18%.",
                              "Items not listed are not quoted.",
                              "Delivery 3 weeks. Payment 30 days."]):
        d.text((90, fy + i * 36), line, font=f("Verdana.ttf", 23), fill=INK)
    d.text((90, fy + 190), "for Tejas Systems", font=f("Verdana.ttf", 24), fill=INK)

    pen = ImageFont.truetype(F + "Bradley Hand Bold.ttf", 40)
    for n, yt, yb, px, pw in rows:
        t = KEY["H4"][n]
        if not t.get("handwritten"):
            continue
        mid = (yt + yb) // 2
        d.line([px - 10, mid + 7, px + pw + 12, mid - 6], fill=PEN, width=5)
        chip = Image.new("RGBA", (300, 90), (0, 0, 0, 0))
        ImageDraw.Draw(chip).text((6, 6), inr(t["price"]), font=pen, fill=PEN + (255,))
        chip = chip.rotate(-3.5, resample=Image.BICUBIC, expand=True)
        img.paste(chip, (int(px - 260), int(yt - 34)), chip)
    return img, rows


def build_h4():
    page, rows = render_h4()
    page = glare(page, rows, target_line=21)
    a, ja = shoot(page, [(240, 250), (2050, 214), (2040, 2840), (220, 2790)],
                  blur=0.8, noise=3.6, jpeg=78)
    a.save(OUT / FILES["H4a"], "JPEG", quality=ja, optimize=True)
    b, jb = shoot(curl(page, 14), [(320, 280), (1990, 220), (2080, 2800), (160, 2880)],
                  blur=1.4, noise=8.0, exposure=0.64, motion=14, warmth=12, jpeg=55)
    b.save(OUT / FILES["H4b"], "JPEG", quality=jb, optimize=True)


# ---------------------------------------------------------------------------
# H5: eml
# ---------------------------------------------------------------------------

def build_h5():
    v = VENDORS["H5"]
    k = KEY["H5"]
    m = EmailMessage()
    m["From"] = f"{v['contact']} <{v['email']}>"
    m["To"] = f"{RFX['buyer_contact']['name']} <{RFX['buyer_contact']['email']}>"
    m["Subject"] = f"Re: {RFX['id']} enquiry - our rates"
    m["Date"] = format_datetime(datetime(2026, 9, 17, 18, 42))
    m["Message-ID"] = "<heldout-h5-0917@saptarishi-infra.example>"
    m.set_content(
        f"Hi {RFX['buyer_contact']['name'].split()[0]},\n\n"
        "Sorry for the delay, was travelling. Our rates for the items we can do,\n"
        "all ex-GST and delivered:\n\n"
        f"  - SFF desktops: {k[4]['price']:,} each\n"
        f"  - 27 inch QHD monitors: {k[6]['price']:,} per unit\n"
        f"  - laptop backpacks: Rs {k[8]['price']:,}\n"
        f"  - Cat6A 2m patch cords: Rs {k[22]['price']:,} per box of 25 (we stock 25s, not 50s)\n"
        f"  - 3kVA rack UPS with NMC: {k[25]['price']:,}\n\n"
        "The access points at the same rate as our last invoice to you. Server\n"
        "racks we cannot supply this time.\n\n"
        "Validity 30 days. Formal quote on letterhead to follow if needed.\n\n"
        f"Regards,\n{v['contact']}\n{v['name']}, {v['city']}\n\n"
        f"On 4 Sep 2026, {RFX['buyer_contact']['name']} wrote:\n"
        f"> Please find attached our enquiry {RFX['id']}. Kindly quote by 18 Sep.\n"
    )
    (OUT / FILES["H5"]).write_bytes(bytes(m))


# ---------------------------------------------------------------------------

def main():
    OUT.mkdir(parents=True, exist_ok=True)
    build_h1()
    build_h2()
    build_h3()
    build_h4()
    build_h5()

    key = {code: {str(n): r for n, r in rows.items()} for code, rows in KEY.items()}
    (OUT / "answer-key.json").write_text(json.dumps(dict(
        note="Held-out answer key. Written from the same table the documents are "
             "rendered from, and committed before any model call.",
        quotes=key), indent=2))

    manifest = [
        dict(label="H1 Kestrel   xlsx", file=FILES["H1"], key="H1", kind="doc"),
        dict(label="H2 Tarang    pdf", file=FILES["H2"], key="H2", kind="doc"),
        dict(label="H3 Nilgiri   docx", file=FILES["H3"], key="H3", kind="doc"),
        dict(label="H4 Tejas photo moderate", file=FILES["H4a"], key="H4", kind="photo"),
        dict(label="H4 Tejas photo low light", file=FILES["H4b"], key="H4", kind="photo"),
        dict(label="H5 Saptarishi eml", file=FILES["H5"], key="H5", kind="doc"),
    ]
    (OUT / "manifest.json").write_text(json.dumps(manifest, indent=2))

    total = 0
    for m in manifest:
        n = len(priced(m["key"]))
        total += n
        print(f"  {m['file']:44} {n:3} priced lines")
    print(f"  {total} priced lines across {len(manifest)} inputs")


if __name__ == "__main__":
    main()
